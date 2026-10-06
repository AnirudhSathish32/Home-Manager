"""The year-end CPA pack: one tax year's records as an Excel workbook to hand an accountant (docs/taxes.md, "CPA pack").

It shows what the app already computes; it does no tax math of its own. Everything is read from one database snapshot,
every foreign amount carries the ECB rate it was converted at, and a Needs review sheet lists what is unfinished. Packs
are kept under <store>/Reports/<year>/, outside the library, so they are never read back in as evidence. A pack is never
overwritten: the same data gives back the same pack, and changed data makes a new file beside the old one.

Python computes every value; the workbook holds values only. Text from documents is written as text, never as a formula.
"""

from datetime import date
from decimal import Decimal
import errno
import hashlib
import json
import os

from ..core.money import EXPONENTS, format_minor
from ..core.paths import safe_path
from ..library.storage import digest_file, now
from .fx import REPORTING, EcbRates, FxError
from .ledger import COUNTABLE, PENDING, RECEIPT_SPENT, SPENDING, STANDALONE_RECEIPT, TRANSACTION_CATEGORY, TRANSACTION_SPENT
from .tax_lots import realized
from .tax_tags import TaxTags

TEMPLATE_VERSION = "cpa-pack-v1"
SHEETS = ("Summary", "Return", "Income", "Write-offs", "Write-off items", "Investments", "Tax forms", "Transactions", "Exchange rates",
          "Needs review", "Manifest")
MAX_ROWS = 100_000  # Per sheet. A bigger year fails clearly instead of being cut short.
EXACT_LIMIT = 2**53  # Larger minor amounts are written as exact text: a spreadsheet number would round them.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


class PackError(RuntimeError):
    """A pack that could not be made, for a reason the user may see. Nothing was saved."""


def M(minor, currency):  # noqa: N802 — a money cell
    return None if minor is None else {"m": int(minor), "c": currency}


def D(day):  # noqa: N802 — a date cell
    return {"d": day} if day else None


def sheet(columns, rows, widths=None):
    if len(rows) > MAX_ROWS:
        raise PackError(f"The year has more than {MAX_ROWS:,} rows for one sheet; the pack was not made.")
    return {"columns": list(columns), "rows": rows, "widths": widths or [18] * len(columns)}


class Collector:
    """Reads one snapshot into the pack's sheets. snap: a store whose connection() is the snapshot's (finance/dashboard.py)."""

    def __init__(self, snap, db, year, tax, today):
        self.snap, self.db, self.year, self.tax, self.today = snap, db, year, tax, today
        self.rates, self.used, self.review = EcbRates(snap), {}, []
        self.start, self.end = f"{year}-01-01", f"{year}-12-31"

    def flag(self, area, item, detail=""):
        self.review.append([area, item, detail])

    def usd(self, minor, currency, day, what):
        """(USD minor, basis, rate handle) for an amount, converted at the ECB rate for its day; flagged when it can't be."""
        if minor is None:
            return None, "", None
        if currency == REPORTING:
            return minor, "usd", None
        try:
            handle = self.rates.lookup(currency, day, self.db)
        except FxError as exc:
            self.flag("Exchange rates", what, f"{format_minor(minor, currency)} on {day} was not converted: {exc}")
            return None, f"unresolved ({exc.code})", None
        assert handle is not None
        self.used[handle.rate_id] = handle
        return handle.convert(minor), "reference_conversion", handle

    def collect(self):
        from .tools import FinanceTools, PeriodInput  # tools imports reconcile, which imports tax_tags
        spending = FinanceTools(self.snap).get_spending(PeriodInput(start=self.start, end=self.end))
        sheets = {"Transactions": self.transactions(spending), "Write-offs": None, "Write-off items": None}
        sheets["Write-offs"], sheets["Write-off items"] = self.write_offs()
        sheets["Return"] = self.return_sheet()
        sheets["Income"] = self.income()
        sheets["Investments"] = self.investments()
        sheets["Tax forms"] = self.tax_forms()
        self.pending()
        sheets["Exchange rates"] = sheet(["Rate id", "Requested for", "ECB date", "USD per unit", "USD per EUR", "Currency per EUR", "Downloaded set", "Set SHA-256"],
                                         [[h.rate_id, D(h.requested_date), D(h.rate_date), f"{h.usd_per_unit:.10f}", h.usd_per_eur, h.currency_per_eur,
                                           h.rate_set_id, h.set_sha256] for h in sorted(self.used.values(), key=lambda h: h.rate_id)],
                                         [22, 14, 12, 16, 12, 16, 14, 66])
        sheets["Needs review"] = sheet(["Area", "Item", "Detail"], self.review, [18, 60, 90])
        sheets["Summary"] = self.summary(spending)
        return {name: sheets[name] for name in SHEETS if name in sheets}

    # Sheets ------------------------------------------------------------------------------

    def summary(self, spending):
        # The return's own input: the gathered values with the typed ones over them (tax_year.merge), so the pack and the return agree.
        estimate, values = self.tax["return"], self.tax["input"]
        result = estimate["result_minor"]
        wages = sum(job["wages"] for job in values["jobs"])
        write_offs = sum(line["counted_minor"] for line in TaxTags(self.snap).year(self.year, REPORTING)["lines"])
        rows = [["Tax year", str(self.year), ""], ["Filing status", self.tax["filing_status_name"], ""],
                ["As of", D(self.today.isoformat()), "Projected to Dec 31: the year is still in progress." if self.year >= self.today.year else ""],
                ["Estimated result", M(abs(result), REPORTING) if result is not None else "Not estimated",
                 "See Needs review" if result is None else "Refund" if result > 0 else "Owed" if result < 0 else "Even"],
                ["Total tax", M(estimate.get("total_tax_minor"), REPORTING), "See Return"],
                ["Total payments", M(estimate.get("payments_minor"), REPORTING), "Withholding, estimated tax and refundable credits"],
                ["Adjusted gross income", M(estimate.get("agi_minor"), REPORTING), ""],
                ["Wages", M(wages, REPORTING), "From confirmed pay stubs and your typed values; see Income"],
                *[[label, M(values.get(key, 0), REPORTING), self.source(key)] for key, label in (
                    ("interest", "Interest"), ("ordinary_dividends", "Ordinary dividends"), ("short_term_gain", "Short-term gains"),
                    ("long_term_gain", "Long-term gains"))],
                ["Write-offs counted", M(write_offs, REPORTING), "Confirmed USD tags; see Write-offs"]]
        usd = spending["usd_total"]
        net_usd = usd["net"]["minor"] if usd else next((row["net_spending"]["minor"] for row in spending["by_currency"] if row["currency"] == REPORTING), 0)
        rows.append(["Household spending (net)", M(net_usd, REPORTING),
                     "All currencies in USD" + ("; some amounts could not be converted, see Needs review" if usd and usd["status"] == "partial" else "")])
        rows += [[f"Spending in {row['currency']}", M(row["net_spending"]["minor"], row["currency"]), "Before conversion"]
                 for row in spending["by_currency"] if row["currency"] != REPORTING]
        complete = estimate["complete"] and not self.review
        rows.append(["Status", "Complete" if complete else "Partial", "Nothing is waiting" if complete else f"{len(self.review)} item(s) on Needs review"])
        return sheet(["Item", "Value", "Note"], rows, [28, 22, 80])

    def source(self, key):
        """Where a return field came from: typed on the Taxes page, else the gathered records' source."""
        if self.tax["inputs"].get("fields", {}).get(key) not in (None, ""):
            return "Typed on the Taxes page"
        return self.tax["gathered"]["sources"].get(key, "")

    def return_sheet(self):
        estimate = self.tax["return"]
        rows = [[line["section"] or "", line["label"], M(line["amount_minor"], REPORTING), line["how"]] for line in estimate["lines"]]
        state = estimate.get("state")
        if state and state.get("state"):
            rows += [["state", f"{state['state']} taxable income", M(state.get("taxable_minor"), REPORTING), ""],
                     ["state", f"{state['state']} tax", M(state.get("tax_minor"), REPORTING), ""],
                     ["state", f"{state['state']} withheld", M(state.get("withheld_minor"), REPORTING), ""]]
        rows += [["note", note, None, ""] for note in estimate["notes"]]
        for key in estimate["missing"]:
            self.flag("Return", f"Missing figure: {key}", "The estimate needs this figure (Taxes → Figures).")
        for code, status in self.tax["tables"].items():
            if status != "verified":
                self.flag("Return", f"{code} tax tables are {status}", "Confirm the year's tables on the Taxes page.")
        return sheet(["Section", "Line", "Amount", "How"], rows, [16, 48, 18, 70])

    def income(self):
        gathered, merged = self.tax["gathered"], self.tax["input"]["jobs"]
        # The return's jobs (tax_year.merge lists the gathered ones first, in order, then the ones typed in), typed values over records.
        rows = []
        for index, job in enumerate(merged):
            found = gathered["jobs"][index] if index < len(gathered["jobs"]) else None
            typed = found and any(job[key] != found["values"].get(key) for key in ("wages", "federal_withheld", "state_withheld"))
            note = (f"{found['stubs']} stubs, last {found['last_pay_date']}, {found['pay_frequency'] or ''}".strip(", ") + ("; typed values over them" if typed else "")
                    if found else "Typed on the Taxes page")
            rows.append(["Job", job["name"], M(job["wages"], REPORTING), M(job["federal_withheld"], REPORTING), M(job["state_withheld"], REPORTING), note])
        rows += [["Field", key, M(self.tax["input"].get(key, value), REPORTING), None, None, self.source(key)]
                 for key, value in sorted(gathered["values"].items()) if isinstance(value, int) and not isinstance(value, bool)]
        return sheet(["Kind", "Name", "Amount", "Federal withheld", "State withheld", "Source"], rows, [8, 34, 16, 16, 16, 70])

    def write_offs(self):
        tags = TaxTags(self.snap)
        counted = tags.year(self.year, REPORTING)
        in_estimate = {tag_id for line in counted["lines"] for tag_id in line["tag_ids"]}
        totals = [[line["kind_label"], line["line_label"], line["business"] or "", line["items"], M(line["amount_minor"], REPORTING),
                   M(line["counted_minor"], REPORTING)] for line in counted["lines"]]
        totals += [["note", note, "", None, None, None] for note in counted["notes"]]
        items = []
        for tag in sorted(tags.list(year=self.year), key=lambda tag: (tag["tax_date"], tag["id"])):
            if tag["review_status"] == "rejected":
                continue
            what = tag["target"].get("description") or f"{tag['target']['type']} {tag['target']['id']}"
            usd, basis, handle = self.usd(tag["counted_minor"], tag["currency"], tag["tax_date"], f"Write-off: {what}")
            if tag["review_status"] != "verified":
                self.flag("Write-offs", f"Unconfirmed tag: {what} on {tag['tax_date']}", "Confirm or reject it on the Taxes page.")
            elif tag["id"] not in in_estimate and tag["currency"] != REPORTING:
                self.flag("Write-offs", f"{what} is in {tag['currency']}", f"The estimate counts USD tags only; about {format_minor(usd, REPORTING) if usd is not None else 'an unconverted amount'}"
                          " at the ECB rate. Enter it under “Typed over the records” or tell your CPA.")
            items.append([D(tag["tax_date"]), tag["kind_label"], tag["line_label"], tag["business"] or "", what, M(tag["amount_minor"], tag["currency"]),
                          M(tag["counted_minor"], tag["currency"]), M(usd, REPORTING), basis, handle.rate_id if handle else "",
                          "Yes" if tag["id"] in in_estimate else "No", tag["review_status"], tag["id"],
                          tag["transaction_id"] or tag["receipt_id"] or tag["receipt_item_id"] or ""])
        return (sheet(["Kind", "Line", "Business", "Items", "Amount", "Counted"], totals, [22, 34, 22, 8, 16, 16]),
                sheet(["Date", "Kind", "Line", "Business", "Item", "Amount", "Counted", "Counted (USD)", "Basis", "Rate id", "In estimate", "Status",
                       "Tag id", "Record id"], items, [12, 20, 28, 18, 40, 16, 16, 16, 20, 22, 11, 11, 8, 10]))

    def investments(self):
        rows = []
        accounts = self.db.execute("SELECT a.id,a.name,a.institution,a.currency,coalesce(a.tax_treatment,k.tax_treatment) AS treatment,k.label "
                                   "FROM investment_accounts a JOIN investment_kinds k ON k.key=a.kind ORDER BY a.name,a.id").fetchall()
        for account in accounts:
            for event in self.db.execute("SELECT event_date,event_type,contribution_source,amount_minor,quantity,note,document_id FROM investment_events "
                                         "WHERE account_id=? AND review_status='verified' AND substr(event_date,1,4)=? ORDER BY event_date,id",
                                         (account["id"], str(self.year))):
                rows.append([account["name"], account["institution"], account["label"], account["treatment"] or "", D(event["event_date"]),
                             event["event_type"], event["contribution_source"] or "", M(event["amount_minor"], account["currency"]), event["quantity"] or "",
                             event["note"], event["document_id"] or ""])
            if account["treatment"] == "taxable":
                gains = realized(self.db, [account["id"]], self.year)
                for label, key in (("Realized proceeds", "proceeds_minor"), ("Realized cost", "cost_minor"), ("Short-term gain", "short_minor"),
                                   ("Long-term gain", "long_minor")):
                    if gains[key]:
                        rows.append([account["name"], account["institution"], account["label"], "taxable", None, label, "",
                                     M(gains[key], account["currency"]), "", "From tax lots", ""])
                for gap in gains["missing"]:
                    self.flag("Investments", f"{account['name']}: {gap['shares']} shares sold on {gap['date']} have no purchase lot",
                              "Add the purchase so the gain can be worked out.")
        return sheet(["Account", "Institution", "Kind", "Tax treatment", "Date", "Activity", "Source", "Amount", "Shares", "Note", "Document"],
                     rows, [24, 20, 18, 14, 12, 18, 10, 16, 10, 40, 10])

    def tax_forms(self):
        rows = []
        for row in self.db.execute("SELECT f.id,f.institution,f.last_four,f.review_status,f.currency,f.document_id,b.form,b.box,b.label,b.amount_minor "
                                   "FROM tax_forms f JOIN tax_form_boxes b ON b.form_id=f.id WHERE f.tax_year=? AND f.review_status<>'rejected' "
                                   "ORDER BY b.form,f.institution,f.id,b.id", (self.year,)):
            rows.append([row["form"], row["institution"], row["last_four"] or "", row["box"], row["label"], M(row["amount_minor"], row["currency"]),
                         row["review_status"], row["document_id"] or ""])
        for (count,) in self.db.execute("SELECT count(*) FROM tax_forms WHERE tax_year=? AND review_status='proposed'", (self.year,)):
            if count:
                self.flag("Tax forms", f"{count} tax form(s) not confirmed", "Confirm them on the Investments page; unconfirmed forms don't count.")
        return sheet(["Form", "Payer", "Last four", "Box", "Label", "Amount", "Status", "Document"], rows, [10, 26, 9, 6, 40, 16, 10, 10])

    def transactions(self, spending):
        """Every counted line and every receipt that still counts on its own, with its USD value and the basis of it. The
        Spent (USD) column adds up to the Summary's household spending, which is checked before the pack is saved."""
        links = {}
        for row in self.db.execute("SELECT l.transaction_id,l.receipt_id,l.estimate_rate_id,l.estimate_minor,r.currency,r.total_minor "
                                   "FROM transaction_receipt_links l JOIN receipts r ON r.id=l.receipt_id WHERE l.review_status<>'rejected' ORDER BY l.id"):
            links.setdefault(row["transaction_id"], dict(row))
        spending_types = (*SPENDING, "refund")
        rows, spent_total = [], 0
        for t in self.db.execute(f"SELECT t.id,t.posted_date,coalesce(t.transaction_date,t.posted_date) AS day,a.display_name AS account,t.description_raw,"
                                 f"t.transaction_type,{TRANSACTION_CATEGORY} AS category,t.amount_minor,t.currency,{TRANSACTION_SPENT} AS spent,t.source_document_id "
                                 f"FROM transactions t JOIN accounts a ON a.id=t.account_id WHERE {COUNTABLE} AND t.posted_date BETWEEN ? AND ? "
                                 "ORDER BY t.posted_date,t.id", (self.start, self.end)).fetchall():
            what = f"{t['description_raw']} on {t['posted_date']}"
            usd, basis, handle = self.usd(t["amount_minor"], t["currency"], t["day"], what)
            spent = None
            if t["transaction_type"] in spending_types:
                net = t["spent"]  # A refund is negative: this person's part of it when shared.
                spent = net if t["currency"] == REPORTING else (handle.convert(net) if handle else None)
                spent_total += spent or 0
            link, note = links.get(t["id"]), ""
            if link and link["currency"] != t["currency"]:
                basis = "actual_settlement"
                note = (f"Receipt {format_minor(link['total_minor'], link['currency'])} matched; this charge counts"
                        + (f" (ECB estimate {format_minor(link['estimate_minor'], REPORTING)}, {link['estimate_rate_id']})" if link["estimate_minor"] is not None else ""))
            rows.append([D(t["posted_date"]), t["account"], t["description_raw"], t["transaction_type"], t["category"], M(t["amount_minor"], t["currency"]),
                         M(usd, REPORTING), M(spent, REPORTING), basis, handle.rate_id if handle else "", link["receipt_id"] if link else "", note,
                         t["id"], t["source_document_id"] or ""])
        for r in self.db.execute(f"SELECT r.id,r.purchase_date,coalesce(m.canonical_name,'Unknown merchant') AS merchant,r.category,r.total_minor,r.currency,"
                                 f"{RECEIPT_SPENT} AS spent,r.document_id FROM receipts r LEFT JOIN merchants m ON m.id=r.merchant_id "
                                 f"WHERE {STANDALONE_RECEIPT} AND r.purchase_date BETWEEN ? AND ? ORDER BY r.purchase_date,r.id", (self.start, self.end)).fetchall():
            usd, basis, handle = self.usd(-r["total_minor"], r["currency"], r["purchase_date"], f"Receipt from {r['merchant']} on {r['purchase_date']}")
            spent = r["spent"] if r["currency"] == REPORTING else (handle.convert(r["spent"]) if handle else None)
            spent_total += spent or 0
            rows.append([D(r["purchase_date"]), "Receipt (no card or bank line yet)", r["merchant"], "receipt", r["category"] or "uncategorized",
                         M(-r["total_minor"], r["currency"]), M(usd, REPORTING), M(spent, REPORTING), basis, handle.rate_id if handle else "", r["id"],
                         "Counts on its own until a statement line replaces it", "", r["document_id"] or ""])
        expected = spending["usd_total"]["net"]["minor"] if spending["usd_total"] else next(
            (row["net_spending"]["minor"] for row in spending["by_currency"] if row["currency"] == REPORTING), 0)
        if spent_total != expected:
            raise PackError("The pack's spending rows do not add up to the household spending total; nothing was saved.")
        return sheet(["Posted", "Account", "Description", "Type", "Category", "Amount", "Amount (USD)", "Spent (USD)", "Basis", "Rate id", "Receipt",
                      "Note", "Transaction id", "Document"], rows, [12, 24, 40, 11, 16, 16, 16, 16, 20, 22, 9, 50, 13, 10])

    def pending(self):
        year = str(self.year)
        checks = [("Receipts", "SELECT count(*) FROM receipts WHERE review_status IN ('proposed','needs_review') AND substr(purchase_date,1,4)=?", (year,),
                   "receipt(s) waiting for review", "Approve or reject them on the Review page."),
                  ("Transactions", f"SELECT count(*) FROM transactions t WHERE {PENDING} AND substr(t.posted_date,1,4)=?", (year,),
                   "statement line(s) not counted yet", "Review or reconcile their statements."),
                  ("Matching", "SELECT count(*) FROM reconciliation_issues WHERE status='open'", (),
                   "open matching question(s)", "Answer them on the Review page; a receipt and its charge could otherwise both count.")]
        for area, sql, params, label, detail in checks:
            count = self.db.execute(sql, params).fetchone()[0]
            if count:
                self.flag(area, f"{count} {label}", detail)


# Workbook ----------------------------------------------------------------------------------

def fingerprint(sheets):
    return hashlib.sha256(json.dumps([TEMPLATE_VERSION, sheets], sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def manifest_sheet(manifest):
    rows = [[key, json.dumps(value, sort_keys=True) if isinstance(value, (dict, list)) else str(value)] for key, value in manifest.items()]
    return sheet(["Field", "Value"], rows, [22, 100])


def render(sheets, stream):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    bold = Font(bold=True)
    for name, data in sheets.items():
        ws = book.create_sheet(name)
        ws.append(data["columns"])
        for cell in ws[1]:
            cell.font = bold
        for values in data["rows"]:
            ws.append([None] * len(values))
            for column, value in enumerate(values, 1):
                write_cell(ws.cell(row=ws.max_row, column=column), value)
        for index, width in enumerate(data["widths"], 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A2"
        if data["rows"]:
            ws.auto_filter.ref = ws.dimensions
    book.save(stream)


def write_cell(cell, value):
    if isinstance(value, dict) and "m" in value:
        exponent = EXPONENTS[value["c"]]
        exact = Decimal(value["m"]).scaleb(-exponent)
        if abs(value["m"]) < EXACT_LIMIT:
            cell.value = exact
            cell.number_format = "#,##0" + ("." + "0" * exponent if exponent else "") + f' "{value["c"]}"'
            return
        value = f"{exact:.{exponent}f} {value['c']}"  # Exact text: too large for a spreadsheet number.
    elif isinstance(value, dict) and "d" in value:
        cell.value = date.fromisoformat(value["d"])
        cell.number_format = "yyyy-mm-dd"
        return
    if value is None or (isinstance(value, int) and not isinstance(value, bool)):
        cell.value = value
        return
    text = str(value)
    cell.value = text
    cell.data_type = "s"  # Never a formula, whatever the text starts with.
    if text.startswith(FORMULA_START):
        cell.quotePrefix = True


def validate(path, sheets):
    """Reopen the saved workbook: the sheets, row counts and Summary values are as written, and no cell is a formula."""
    from openpyxl import load_workbook
    with open(path, "rb") as stream:  # By handle: openpyxl refuses a name that doesn't end in .xlsx, such as the .partial.
        try:
            book = load_workbook(stream, read_only=True)
        except Exception as exc:  # A workbook that does not reopen is not published.
            raise PackError("The saved workbook could not be reopened; nothing was saved.") from exc
        check(book, sheets)


def check(book, sheets):
    try:
        if book.sheetnames != list(sheets):
            raise PackError("The saved workbook is missing sheets; nothing was saved.")
        for name, data in sheets.items():
            rows = list(book[name].iter_rows())
            if len(rows) != len(data["rows"]) + 1 or any(cell.data_type == "f" for row in rows for cell in row):
                raise PackError(f"The saved workbook's {name} sheet is not as written; nothing was saved.")
        for written, row in zip(sheets["Summary"]["rows"], list(book["Summary"].iter_rows(values_only=True))[1:]):
            value = written[1]
            if isinstance(value, dict) and "m" in value and abs(value["m"]) < EXACT_LIMIT:
                if Decimal(str(row[1])) != Decimal(value["m"]).scaleb(-EXPONENTS[value["c"]]):
                    raise PackError("The saved workbook's totals differ from the computed ones; nothing was saved.")
    finally:
        book.close()


# Packs ------------------------------------------------------------------------------------

class CpaPacks:
    def __init__(self, store):
        self.store = store
        self.root = safe_path(store.root / "Reports")

    def build(self, year, tax_for, today=None):
        """tax_for(snapshot_store) gives the year's tax view (Manager.tax_view) from the same snapshot."""
        from .dashboard import SnapshotStore
        today = today or date.today()
        if not 2000 <= year <= today.year:
            raise ValueError("Choose a tax year up to this one.")
        with self.store.connection() as db:
            db.execute("BEGIN")
            snap = SnapshotStore(self.store, db)
            collector = Collector(snap, db, year, tax_for(snap), today)
            sheets = collector.collect()
            schema = db.execute("PRAGMA user_version").fetchone()[0]
            sets = sorted({(h.rate_set_id, h.set_sha256) for h in collector.used.values()})
        key = fingerprint(sheets)
        request_key = f"cpa_pack:{year}:{key}"
        with self.store.connection() as db:
            existing = db.execute("SELECT * FROM generated_reports WHERE request_key=?", (request_key,)).fetchone()
        if existing:
            path = self.file(existing)
            if path.exists() and digest_file(path) == existing["sha256"]:
                return {**self.view(existing), "reused": True}
        manifest = {"template_version": TEMPLATE_VERSION, "generated_at": now(), "tax_year": year, "as_of": today.isoformat(), "schema_version": schema,
                    "data_fingerprint": key, "rate_sets": [{"id": set_id, "sha256": sha} for set_id, sha in sets],
                    "rows": {name: len(data["rows"]) for name, data in sheets.items()}, "needs_review": len(collector.review),
                    "source": "Home Manager records; values computed by the app, not by spreadsheet formulas."}
        sheets["Manifest"] = manifest_sheet(manifest)
        folder = safe_path(self.root / str(year))
        final = safe_path(folder / f"cpa-pack_{year}__{key[:12]}.xlsx")
        partial = safe_path(final.with_name(final.name + ".partial"))
        try:
            folder.mkdir(parents=True, exist_ok=True)
            partial.unlink(missing_ok=True)  # Only this pack's own unfinished file, from an attempt that was cut short.
            with open(partial, "xb") as stream:
                render(sheets, stream)
                stream.flush()
                os.fsync(stream.fileno())
            validate(partial, sheets)
            try:
                os.link(partial, final)  # Never replaces a file already there.
            except FileExistsError:
                if not existing:
                    raise PackError("A different file already has this pack's name; nothing was saved.") from None
                raise PackError("The saved pack for this data was changed or removed from Reports. Move the old file away, then build again.") from None
        except PermissionError as exc:
            raise PackError("The Reports folder or the pack file is locked (is it open in Excel?). Close it and try again.") from exc
        except OSError as exc:
            if exc.errno == errno.ENOSPC:
                raise PackError("There is not enough disk space to save the CPA pack.") from exc
            raise
        finally:
            partial.unlink(missing_ok=True)
        sha = digest_file(final)
        with self.store.connection() as db:
            pack_id = db.execute("INSERT INTO generated_reports(kind,request_key,relative_path,sha256,byte_size,tax_year,created_at,manifest_json) "
                                 "VALUES('cpa_pack',?,?,?,?,?,?,?)", (request_key, final.relative_to(self.store.root).as_posix(), sha,
                                                                      final.stat().st_size, year, manifest["generated_at"], json.dumps(manifest, sort_keys=True))).lastrowid
            row = db.execute("SELECT * FROM generated_reports WHERE id=?", (pack_id,)).fetchone()
        return {**self.view(row), "reused": False}

    def list(self, year=None):
        with self.store.connection() as db:
            rows = db.execute("SELECT * FROM generated_reports WHERE kind='cpa_pack'" + (" AND tax_year=?" if year else "") + " ORDER BY id DESC",
                              (year,) if year else ()).fetchall()
        return [self.view(row) for row in rows]

    def path(self, pack_id):
        """The pack's file, checked against the hash recorded when it was made."""
        with self.store.connection() as db:
            row = db.execute("SELECT * FROM generated_reports WHERE id=? AND kind='cpa_pack'", (pack_id,)).fetchone()
        if row is None:
            raise ValueError("CPA pack not found.")
        path = self.file(row)
        if not path.exists() or digest_file(path) != row["sha256"]:
            raise PackError("The pack file was changed or removed from Reports since it was made.")
        return path, path.name

    def file(self, row):
        return safe_path(self.store.root.joinpath(*row["relative_path"].split("/")))

    @staticmethod
    def view(row):
        manifest = json.loads(row["manifest_json"])
        return {"id": row["id"], "tax_year": row["tax_year"], "name": row["relative_path"].rsplit("/", 1)[-1], "created_at": row["created_at"],
                "sha256": row["sha256"], "byte_size": row["byte_size"], "as_of": manifest.get("as_of"), "needs_review": manifest.get("needs_review", 0),
                "rows": manifest.get("rows", {})}
