"""The year-end CPA pack (finance/cpa_pack.py, docs/taxes.md "CPA pack"), with synthetic records and a synthetic ECB file only."""

from decimal import Decimal
import errno
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from conftest import documents_by_name, inbox_scan
from home_manager.app.api import create_app
from home_manager.finance import cpa_pack
from home_manager.finance.ledger import Ledger
from home_manager.finance.tax_tags import TaxTags
from home_manager.library.scanner import ScanLimits
from test_fx import FakeEcb

INJECTED = '=HYPERLINK("https://example.invalid","click")'


@pytest.fixture
def app_books(tmp_path):
    """A card used at home and in Mexico in March 2025: USD charges, a charity gift, and receipts in MXN and PHP."""
    app = create_app(tmp_path / "control", "t", limits=ScanLimits(stability_seconds=0))
    with TestClient(app, base_url="http://127.0.0.1:8765", headers={"Authorization": "Bearer t"}) as client:
        client.put("/api/settings", json={"managed_directory": str(tmp_path / "managed")})
        manager = app.state.manager
        manager.rate_fetch = FakeEcb()
        store = manager.store
        inbox_scan(store, {"card.csv": b"date,amount\n", "oxxo.png": b"synthetic mxn receipt", "jollibee.png": b"synthetic php receipt"})
        docs = documents_by_name(store)
        ledger = Ledger(store)
        card = ledger.create_account("Travel Card", "credit_card", "USD")

        def charge(*rows):
            with store.connection() as db:
                return ledger.insert_transactions(db, card, [{"posted_date": day, "description": text, "amount_minor": amount, "currency": "USD",
                                                              "locator": {"rows": [index]}} for index, (day, text, amount) in enumerate(rows, 1)],
                                                  "import", {"document_id": docs["card.csv"]["id"], "blob_hash": docs["card.csv"]["current_hash"],
                                                             "source_key": f"import:{rows[0][1]}"})[0]

        def receipt(name, merchant, day, total, currency):
            doc = docs[name]
            return ledger.publish_receipt({"merchant": merchant, "purchase_date": day, "subtotal_minor": None, "tax_minor": None, "tip_minor": None,
                                           "total_minor": total, "currency": currency, "issues": [], "items": [], "locator": {"line_ids": ["line-1"]}},
                                          {"document_id": doc["id"], "blob_hash": doc["current_hash"], "source_key": f"extraction:{name}", "run_id": name},
                                          "verified")["id"]

        [gift, _, _] = charge(("2025-03-05", "RED CROSS", -25_000), ("2025-03-06", INJECTED, -100), ("2025-03-20", "GROCER", -5_150))
        oxxo = receipt("oxxo.png", "OXXO", "2025-03-15", 100_000, "MXN")
        receipt("jollibee.png", "Jollibee", "2025-03-13", 30_000, "PHP")  # No PHP rate within seven days: unresolved
        tags = TaxTags(store)
        tags.tag("transaction", gift, "itemized", "charity_cash")
        tags.tag("receipt", oxxo, "itemized", "charity_cash")
        assert client.post("/api/rates/refresh").json()["rates"] > 0
        yield client, manager, charge


def build(client, year=2025):
    response = client.post(f"/api/tax/cpa-packs/{year}")
    assert response.status_code == 201, response.text
    return response.json()


def workbook(client, pack):
    response = client.get(f"/api/tax/cpa-packs/{pack['id']}/file")
    assert response.status_code == 200 and f'filename="{pack["name"]}"' in response.headers["content-disposition"]
    return load_workbook(io.BytesIO(response.content))


def rows(book, name):
    return list(book[name].iter_rows(values_only=True))


def test_the_pack_holds_the_year_with_usd_values_rates_and_what_needs_review(app_books):
    client, manager, _ = app_books
    pack = build(client)
    assert pack["name"].startswith("cpa-pack_2025__") and pack["reused"] is False and pack["needs_review"] > 0
    book = workbook(client, pack)
    assert book.sheetnames == list(cpa_pack.SHEETS)
    summary = {row[0]: row[1:] for row in rows(book, "Summary")[1:]}
    # 250.00 + 1.00 + 51.50 in USD, and 1,000.00 MXN at 1.09/21.80 = 50.00. The PHP receipt has no rate and is left out.
    assert summary["Household spending (net)"][0] == Decimal("352.50")
    assert summary["Spending in MXN"][0] == Decimal("1000.00") and summary["Status"][0] == "Partial"
    assert summary["Write-offs counted"][0] == Decimal("250.00")  # The USD gift; the MXN one is flagged, not counted
    transactions = rows(book, "Transactions")
    injected = next(row for row in transactions if row[2] == INJECTED)
    assert injected[5] == Decimal("-1.00")
    cell = next(cell for cell in book["Transactions"]["C"] if cell.value == INJECTED)
    assert cell.data_type == "s" and cell.quotePrefix
    mxn = next(row for row in transactions if row[2] == "OXXO")
    assert (mxn[5], mxn[6], mxn[7], mxn[8], mxn[9]) == (Decimal("-1000.00"), Decimal("-50.00"), Decimal("50.00"), "reference_conversion", "ecb:2025-03-14:MXN")
    assert [row[0] for row in rows(book, "Exchange rates")[1:]] == ["ecb:2025-03-14:MXN"]
    review = rows(book, "Needs review")
    assert any(row[0] == "Exchange rates" and "Jollibee" in row[1] for row in review)
    assert any(row[0] == "Write-offs" and "MXN" in row[1] for row in review)
    items = rows(book, "Write-off items")
    assert sorted((row[4], row[10]) for row in items[1:]) == [("OXXO", "No"), ("RED CROSS", "Yes")]
    manifest = dict(rows(book, "Manifest")[1:])
    assert manifest["tax_year"] == "2025" and manifest["template_version"] == cpa_pack.TEMPLATE_VERSION
    assert client.get("/api/tax/cpa-packs", params={"year": 2025}).json()["packs"][0]["id"] == pack["id"]


def test_same_data_gives_back_the_same_pack_and_new_data_never_overwrites(app_books):
    client, manager, charge = app_books
    first = build(client)
    again = build(client)
    assert (again["id"], again["sha256"], again["reused"]) == (first["id"], first["sha256"], True)
    charge(("2025-04-02", "BOOKSHOP", -2_000))
    newer = build(client)
    assert newer["id"] != first["id"] and newer["name"] != first["name"]
    folder = manager.store.root / "Reports" / "2025"
    assert sorted(path.name for path in folder.iterdir()) == sorted([first["name"], newer["name"]])
    assert workbook(client, first) and workbook(client, newer)  # Both still served, each checked against its hash
    (folder / first["name"]).write_bytes(b"edited")
    assert client.get(f"/api/tax/cpa-packs/{first['id']}/file").status_code == 409


@pytest.mark.parametrize("error,message", [(OSError(errno.ENOSPC, "No space left on device"), "not enough disk space"),
                                           (PermissionError(13, "locked"), "locked")])
def test_disk_full_or_locked_files_fail_clearly_and_leave_nothing(app_books, monkeypatch, error, message):
    client, manager, _ = app_books

    def broken(sheets, stream):
        stream.write(b"partial")
        raise error
    monkeypatch.setattr(cpa_pack, "render", broken)
    response = client.post("/api/tax/cpa-packs/2025")
    assert response.status_code == 409 and message in response.json()["detail"]
    folder = manager.store.root / "Reports" / "2025"
    assert list(folder.iterdir()) == [] and client.get("/api/tax/cpa-packs").json() == {"packs": []}


def test_endpoints_check_auth_years_and_ids(app_books):
    client, _, _ = app_books
    assert client.post("/api/tax/cpa-packs/1800").status_code == 400
    assert client.post("/api/tax/cpa-packs/2100").status_code == 400  # Not a year that has happened
    assert client.post("/api/tax/cpa-packs/next").status_code == 422
    assert client.get("/api/tax/cpa-packs/999/file").status_code == 400
    assert client.post("/api/tax/cpa-packs/2025", headers={"Authorization": "Bearer wrong"}).status_code == 401
